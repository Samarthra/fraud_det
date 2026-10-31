"""Read CSV, Excel or JSON into (row number, raw values) pairs."""

import csv
import io
import json
import re
from datetime import date, datetime
from pathlib import PurePath
from typing import Any

from openpyxl import load_workbook

OVERFLOW = "__overflow__"


class FileRejected(Exception):
    """The whole file is unusable, as opposed to individual bad rows."""


def normalise_header(name: Any) -> str:
    return re.sub(r"[\s-]+", "_", str(name or "").strip().lower())


def parse_file(filename: str, content: bytes) -> list[tuple[int, Any]]:
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".csv":
        return _parse_csv(content)
    if suffix == ".xlsx":
        return _parse_xlsx(content)
    if suffix == ".json":
        return _parse_json(content)
    raise FileRejected(f"Unsupported file type '{suffix}'. Use .csv, .xlsx or .json.")


def _parse_csv(content: bytes) -> list[tuple[int, Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as e:
        raise FileRejected("CSV file is not valid UTF-8.") from e
    reader = csv.reader(io.StringIO(text, newline=""))
    try:
        header = next(reader, None)
        if header is None:
            raise FileRejected("File is empty.")
        return [
            (reader.line_num, _zip_row(header, cells))
            for cells in reader
        ]
    except csv.Error as e:
        raise FileRejected(f"CSV file could not be read: {e}") from e


def _parse_xlsx(content: bytes) -> list[tuple[int, Any]]:
    try:
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    except Exception as e:
        raise FileRejected("Excel file could not be read.") from e
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise FileRejected("File is empty.")
    return [(i, _zip_row(header, cells)) for i, cells in enumerate(rows, start=2)]


def _parse_json(content: bytes) -> list[tuple[int, Any]]:
    try:
        data = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as e:
        raise FileRejected("JSON file could not be read.") from e
    if isinstance(data, dict) and isinstance(data.get("rows"), list):
        data = data["rows"]
    if not isinstance(data, list):
        raise FileRejected('JSON must be a list of records, or {"rows": [...]}.')
    return [
        (i, {normalise_header(k): v for k, v in item.items()} if isinstance(item, dict) else item)
        for i, item in enumerate(data, start=1)
    ]


def _zip_row(header: tuple | list, cells: tuple | list) -> dict[str, Any]:
    names = [normalise_header(h) for h in header]
    row = {name: _cell(cell) for name, cell in zip(names, cells) if name}
    extra = [c for c in cells[len(names):] if c not in (None, "")]
    if extra:
        row[OVERFLOW] = len(extra)
    return row


def _cell(value: Any) -> Any:
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value
